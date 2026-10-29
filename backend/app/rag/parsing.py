"""Turn uploaded files into normalized Markdown-ish text plus document metadata."""

import io
import re
from dataclasses import dataclass, field
from pathlib import PurePath

SUPPORTED_EXTENSIONS = {".md", ".markdown", ".txt", ".pdf", ".docx"}
DEFAULT_ACL = ("all-employees",)

_FRONT_MATTER = re.compile(r"\A---\s*\n(.*?)\n---\s*\n", re.DOTALL)
_H1 = re.compile(r"^#\s+(.+?)\s*$", re.MULTILINE)


class UnsupportedFileType(ValueError):
    pass


@dataclass(frozen=True)
class ParsedDocument:
    doc_key: str
    title: str
    body: str
    source_path: str
    doc_type: str | None = None
    owner: str | None = None
    version: str | None = None
    updated: str | None = None
    acl: tuple[str, ...] = field(default=DEFAULT_ACL)


def parse_front_matter(text: str) -> tuple[dict[str, str], str]:
    """Split simple `key: value` front matter (between `---` lines) from the body."""
    match = _FRONT_MATTER.match(text)
    if not match:
        return {}, text
    meta: dict[str, str] = {}
    for line in match.group(1).splitlines():
        if ":" not in line or line.lstrip().startswith("#"):
            continue
        key, value = line.split(":", 1)
        meta[key.strip().lower()] = value.strip().strip('"').strip("'")
    return meta, text[match.end() :]


def normalize_text(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace(" ", " ")
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _pdf_to_text(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    pages = [(page.extract_text() or "").strip() for page in reader.pages]
    return "\n\n".join(p for p in pages if p)


def _docx_to_markdown(data: bytes) -> str:
    import docx

    document = docx.Document(io.BytesIO(data))
    lines: list[str] = []
    for para in document.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style = (para.style.name or "").lower() if para.style is not None else ""
        if style == "title":
            lines.append(f"# {text}")
        elif style.startswith("heading"):
            level = int(style.split()[-1]) if style.split()[-1].isdigit() else 2
            lines.append(f"{'#' * min(level + 1, 6)} {text}")
        elif "list" in style:
            lines.append(f"- {text}")
        else:
            lines.append(text)
    return "\n\n".join(lines)


def extract_text(filename: str, data: bytes) -> str:
    ext = PurePath(filename).suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise UnsupportedFileType(
            f"Unsupported file type '{ext}'. Supported: {', '.join(sorted(SUPPORTED_EXTENSIONS))}"
        )
    if ext == ".pdf":
        return _pdf_to_text(data)
    if ext == ".docx":
        return _docx_to_markdown(data)
    return data.decode("utf-8-sig", errors="replace")


def parse_document(
    filename: str,
    data: bytes,
    *,
    source_path: str | None = None,
    acl: tuple[str, ...] | None = None,
) -> ParsedDocument:
    """Parse a file. Front-matter values win; `acl` is a fallback when none is declared."""
    meta, body = parse_front_matter(extract_text(filename, data))
    body = normalize_text(body)
    if not body:
        raise ValueError("Document has no extractable text")

    stem = PurePath(filename).stem
    h1 = _H1.search(body)
    title = meta.get("title") or (h1.group(1) if h1 else stem.replace("-", " ").title())
    declared_acl = tuple(g.strip() for g in meta.get("acl", "").split(",") if g.strip())

    return ParsedDocument(
        doc_key=meta.get("doc_id") or stem,
        title=title,
        body=body,
        source_path=source_path or filename,
        doc_type=meta.get("doc_type"),
        owner=meta.get("owner"),
        version=meta.get("version"),
        updated=meta.get("updated"),
        acl=declared_acl or acl or DEFAULT_ACL,
    )
