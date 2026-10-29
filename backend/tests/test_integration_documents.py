"""Documents API against real Postgres + pgvector."""

import io

import docx
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import select

from app.core.config import Settings
from app.db.models import Chunk, Document
from app.main import create_app
from app.rag.embeddings import HashingEmbedder
from app.rag.ingestion import ingest_directory
from tests.conftest import CORPUS_DIR

pytestmark = pytest.mark.integration


@pytest.fixture
async def corpus(pg_session):
    await ingest_directory(pg_session, CORPUS_DIR, HashingEmbedder())
    return pg_session


def client_for(url: str, groups: list[str] | None = None) -> TestClient:
    settings = Settings(
        _env_file=None,
        llm_provider="fake",
        embedding_provider="hashing",
        database_url=url,
        user_groups=groups or ["all-employees"],
    )
    return TestClient(create_app(settings))


async def test_list_hides_documents_outside_callers_groups(corpus, migrated_db):
    with client_for(migrated_db) as client:
        keys = {d["doc_key"] for d in client.get("/api/v1/documents").json()["documents"]}
    assert len(keys) == 10
    assert not keys & {"HR-007", "LDR-014"}

    with client_for(migrated_db, ["all-employees", "leadership"]) as client:
        keys = {d["doc_key"] for d in client.get("/api/v1/documents").json()["documents"]}
    assert "LDR-014" in keys and "HR-007" not in keys


async def test_upload_markdown_then_it_is_searchable(corpus, migrated_db):
    content = b"# Parking Policy\n\nFree parking at the Austin office garage, level 3."
    with client_for(migrated_db) as client:
        resp = client.post(
            "/api/v1/documents", files={"file": ("parking.md", content, "text/markdown")}
        )
        assert resp.status_code == 201, resp.text
        body = resp.json()
        assert body["status"] == "created"
        assert body["document"]["acl"] == ["all-employees"]
        assert body["document"]["chunk_count"] == 1

        again = client.post("/api/v1/documents", files={"file": ("parking.md", content)})
        assert again.json()["status"] == "unchanged"

        hits = client.post("/api/v1/search", json={"query": "parking in Austin"}).json()
    assert hits["hits"][0]["doc_key"] == "parking"


async def test_upload_docx_keeps_headings(corpus, migrated_db):
    d = docx.Document()
    d.add_heading("Desk Booking", level=0)
    d.add_heading("How to book", level=1)
    d.add_paragraph("Book hot desks in the workplace app up to 14 days ahead. " * 3)
    buf = io.BytesIO()
    d.save(buf)

    with client_for(migrated_db) as client:
        resp = client.post("/api/v1/documents", files={"file": ("desks.docx", buf.getvalue())})
    assert resp.status_code == 201, resp.text
    chunk = await corpus.scalar(select(Chunk).join(Document).where(Document.doc_key == "desks"))
    assert chunk.section == "Desk Booking › How to book"


async def test_upload_rejects_acl_outside_callers_groups(corpus, migrated_db):
    with client_for(migrated_db) as client:
        resp = client.post(
            "/api/v1/documents",
            files={"file": ("secret.md", b"# Secret\n\nSome secret plans for the board meeting.")},
            data={"acl": "leadership"},
        )
    assert resp.status_code == 403


async def test_upload_cannot_overwrite_a_document_the_caller_cannot_see(corpus, migrated_db):
    spoof = b"---\ndoc_id: HR-007\nacl: all-employees\n---\n# Bands\n\nEveryone earns the same now."
    with client_for(migrated_db) as client:
        resp = client.post("/api/v1/documents", files={"file": ("bands.md", spoof)})
    assert resp.status_code == 403
    stored = await corpus.scalar(select(Document).where(Document.doc_key == "HR-007"))
    assert stored.acl == ["hr-confidential"]


async def test_upload_rejects_unsupported_and_empty_files(corpus, migrated_db):
    with client_for(migrated_db) as client:
        assert (
            client.post("/api/v1/documents", files={"file": ("x.png", b"\x89PNG")}).status_code
            == 415
        )
        assert client.post("/api/v1/documents", files={"file": ("e.md", b"   ")}).status_code == 422


async def test_chunk_lookup_enforces_acl(corpus, migrated_db):
    restricted = await corpus.scalar(
        select(Chunk.id).join(Document).where(Document.doc_key == "HR-007")
    )
    public = await corpus.scalar(
        select(Chunk.id).join(Document).where(Document.doc_key == "HR-001")
    )
    with client_for(migrated_db) as client:
        assert client.get(f"/api/v1/chunks/{restricted}").status_code == 404
        ok = client.get(f"/api/v1/chunks/{public}")
    assert ok.status_code == 200
    assert ok.json()["document"]["doc_key"] == "HR-001"
